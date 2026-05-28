"""
模板生成脚本 - 一次性运行，生成所有 docx 模板文件

运行: python scripts/create_templates.py
"""

import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def create_car_purchase_contract():
    """购车合同模板"""
    doc = Document()

    # 标题
    title = doc.add_heading("购车合同", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f"合同编号：{{{{contract_no}}}}")
    doc.add_paragraph("")

    # 甲方乙方
    doc.add_paragraph("甲方（卖方）：{{dealer_name}}")
    doc.add_paragraph("地址：{{dealer_address}}")
    doc.add_paragraph("联系电话：{{dealer_phone}}")
    doc.add_paragraph("")
    doc.add_paragraph("乙方（买方）：{{buyer_name}}")
    doc.add_paragraph("身份证号：{{buyer_id_card}}")
    doc.add_paragraph("联系电话：{{buyer_phone}}")
    doc.add_paragraph("")

    doc.add_heading("第一条 车辆信息", level=1)
    table = doc.add_table(rows=6, cols=2, style="Table Grid")
    cells = [
        ("车辆品牌", "{{car_brand}}"),
        ("车型名称", "{{car_model}}"),
        ("车辆颜色", "{{car_color}}"),
        ("车架号（VIN）", "{{vin}}"),
        ("发动机号", "{{engine_no}}"),
        ("车辆价格（元）", "{{car_price}}"),
    ]
    for i, (label, value) in enumerate(cells):
        table.cell(i, 0).text = label
        table.cell(i, 1).text = value

    doc.add_paragraph("")
    doc.add_heading("第二条 付款方式", level=1)
    doc.add_paragraph("付款方式：{{payment_method}}")
    doc.add_paragraph("首付金额：{{down_payment}} 元")
    doc.add_paragraph("贷款金额：{{loan_amount}} 元")
    doc.add_paragraph("")

    doc.add_heading("第三条 交车约定", level=1)
    doc.add_paragraph("预计交车日期：{{delivery_date}}")
    doc.add_paragraph("交车地点：{{delivery_location}}")
    doc.add_paragraph("")

    doc.add_heading("第四条 双方权利义务", level=1)
    doc.add_paragraph("1. 甲方保证所售车辆为全新正品，符合国家质量标准。")
    doc.add_paragraph("2. 甲方负责办理车辆上牌手续（费用由乙方承担）。")
    doc.add_paragraph("3. 乙方应按约定时间支付车款。")
    doc.add_paragraph("4. 乙方应在交车后 7 日内完成验收。")
    doc.add_paragraph("")

    doc.add_heading("第五条 违约责任", level=1)
    doc.add_paragraph("1. 甲方逾期交车，每逾期一日按车价 0.05% 支付违约金。")
    doc.add_paragraph("2. 乙方逾期付款，每逾期一日按欠款 0.05% 支付违约金。")
    doc.add_paragraph("")

    doc.add_heading("第六条 争议解决", level=1)
    doc.add_paragraph("本合同发生争议，双方协商解决；协商不成，提交甲方所在地人民法院诉讼。")
    doc.add_paragraph("")
    doc.add_paragraph("")

    # 签章区
    doc.add_paragraph("甲方（盖章）：                    乙方（签字）：")
    doc.add_paragraph("")
    doc.add_paragraph("日期：{{sign_date}}                日期：{{sign_date}}")

    return doc


def create_order_agreement():
    """订车协议模板"""
    doc = Document()

    title = doc.add_heading("订车协议", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f"协议编号：{{{{agreement_no}}}}")
    doc.add_paragraph("")

    doc.add_paragraph("甲方（经销商）：{{dealer_name}}")
    doc.add_paragraph("乙方（客户）：{{buyer_name}}")
    doc.add_paragraph("联系电话：{{buyer_phone}}")
    doc.add_paragraph("")

    doc.add_heading("一、预订车辆信息", level=1)
    table = doc.add_table(rows=5, cols=2, style="Table Grid")
    cells = [
        ("车型", "{{car_model}}"),
        ("配置/版本", "{{car_config}}"),
        ("颜色", "{{car_color}}"),
        ("预估价格（元）", "{{estimated_price}}"),
        ("预计到车日期", "{{eta_date}}"),
    ]
    for i, (label, value) in enumerate(cells):
        table.cell(i, 0).text = label
        table.cell(i, 1).text = value

    doc.add_paragraph("")
    doc.add_heading("二、订金条款", level=1)
    doc.add_paragraph("订金金额：{{deposit_amount}} 元")
    doc.add_paragraph("支付方式：{{deposit_payment_method}}")
    doc.add_paragraph("1. 乙方支付订金后，甲方保留车辆名额。")
    doc.add_paragraph("2. 乙方取消订单，订金不予退还。")
    doc.add_paragraph("3. 甲方无法按时交车，双倍返还订金。")
    doc.add_paragraph("")

    doc.add_heading("三、购车流程", level=1)
    doc.add_paragraph("1. 乙方支付订金，签订本协议。")
    doc.add_paragraph("2. 车辆到店后，甲方通知乙方验车。")
    doc.add_paragraph("3. 乙方验车满意后，签订正式购车合同并支付尾款。")
    doc.add_paragraph("")

    doc.add_paragraph("甲方（盖章）：                    乙方（签字）：")
    doc.add_paragraph("")
    doc.add_paragraph("日期：{{sign_date}}                日期：{{sign_date}}")

    return doc


def create_finance_installment_contract():
    """金融分期合同模板"""
    doc = Document()

    title = doc.add_heading("汽车金融分期付款合同", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f"合同编号：{{{{contract_no}}}}")
    doc.add_paragraph("")

    doc.add_paragraph("甲方（贷款人/金融机构）：{{finance_company}}")
    doc.add_paragraph("乙方（借款人）：{{borrower_name}}")
    doc.add_paragraph("身份证号：{{borrower_id_card}}")
    doc.add_paragraph("联系电话：{{borrower_phone}}")
    doc.add_paragraph("")

    doc.add_heading("第一条 贷款信息", level=1)
    table = doc.add_table(rows=7, cols=2, style="Table Grid")
    cells = [
        ("车辆信息", "{{car_model}} ({{car_color}})"),
        ("车辆总价（元）", "{{total_price}}"),
        ("首付金额（元）", "{{down_payment}}"),
        ("贷款金额（元）", "{{loan_amount}}"),
        ("贷款期数", "{{installments}} 期"),
        ("年利率", "{{annual_rate}}%"),
        ("月供金额（元）", "{{monthly_payment}}"),
    ]
    for i, (label, value) in enumerate(cells):
        table.cell(i, 0).text = label
        table.cell(i, 1).text = value

    doc.add_paragraph("")
    doc.add_heading("第二条 还款方式", level=1)
    doc.add_paragraph("还款方式：等额本息")
    doc.add_paragraph("还款日：每月 {{repayment_day}} 日")
    doc.add_paragraph("还款账户：{{repayment_account}}")
    doc.add_paragraph("")

    doc.add_heading("第三条 车辆抵押", level=1)
    doc.add_paragraph("1. 乙方将所购车辆抵押给甲方作为贷款担保。")
    doc.add_paragraph("2. 贷款还清后，甲方协助办理抵押解除手续。")
    doc.add_paragraph("3. 贷款期间，乙方不得转让、出售抵押车辆。")
    doc.add_paragraph("")

    doc.add_heading("第四条 提前还款", level=1)
    doc.add_paragraph("1. 乙方可申请提前还款，需提前 30 日书面通知甲方。")
    doc.add_paragraph("2. 提前还款不收取违约金（满 12 期后）。")
    doc.add_paragraph("")

    doc.add_heading("第五条 逾期处理", level=1)
    doc.add_paragraph("1. 逾期还款，按日加收逾期金额 0.05% 的罚息。")
    doc.add_paragraph("2. 连续逾期 3 期以上，甲方有权收回车辆。")
    doc.add_paragraph("")

    doc.add_paragraph("甲方（盖章）：                    乙方（签字）：")
    doc.add_paragraph("")
    doc.add_paragraph("日期：{{sign_date}}                日期：{{sign_date}}")

    return doc


def main():
    """生成所有模板文件"""
    base_dir = os.path.join(os.path.dirname(__file__), "..", "data", "templates", "contracts")
    os.makedirs(base_dir, exist_ok=True)

    templates = {
        "car_purchase.docx": create_car_purchase_contract,
        "order_agreement.docx": create_order_agreement,
        "finance_installment.docx": create_finance_installment_contract,
    }

    for filename, creator in templates.items():
        filepath = os.path.join(base_dir, filename)
        doc = creator()
        doc.save(filepath)
        print(f"[OK] 已生成: {filepath}")

    # 提案书模板
    proposal_dir = os.path.join(os.path.dirname(__file__), "..", "data", "templates", "proposals")
    os.makedirs(proposal_dir, exist_ok=True)
    doc = Document()
    title = doc.add_heading("购车方案提案书", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"提案编号：{{{{proposal_no}}}}")
    doc.add_paragraph("致：{{customer_name}}")
    doc.add_paragraph("日期：{{proposal_date}}")
    doc.add_paragraph("")
    doc.add_heading("一、客户需求分析", level=1)
    doc.add_paragraph("{{needs_analysis}}")
    doc.add_paragraph("")
    doc.add_heading("二、推荐车型", level=1)
    table = doc.add_table(rows=1, cols=5, style="Table Grid")
    headers = ["车型", "配置", "官方指导价", "优惠方案", "推荐理由"]
    for i, h in enumerate(headers):
        table.cell(0, i).text = h
    doc.add_paragraph("")
    doc.add_heading("三、金融方案", level=1)
    doc.add_paragraph("{{finance_plan}}")
    doc.add_paragraph("")
    doc.add_heading("四、增值服务", level=1)
    doc.add_paragraph("{{value_added_services}}")
    doc.add_paragraph("")
    doc.add_heading("五、购车总费用明细", level=1)
    doc.add_paragraph("{{cost_breakdown}}")
    doc.add_paragraph("")
    doc.add_paragraph("专属顾问：{{sales_rep}}")
    doc.add_paragraph("联系方式：{{sales_phone}}")
    doc.save(os.path.join(proposal_dir, "proposal.docx"))
    print(f"[OK] 已生成: {os.path.join(proposal_dir, 'proposal.docx')}")

    print("\n所有模板生成完成！")


if __name__ == "__main__":
    main()
